"""Montreve development budget, break-even and pro forma.

One table of inputs (low / base / high, each with its source) drives everything:
  - data/model.json: every number the site plan and the presentation quote
  - out/Montreve_Budget_ProForma.xlsx: the same model with live formulas (change an input, the workbook follows)

Three programs are priced on the same 24 for-sale cabins (the 7 financed homes + 17 new cabins on the remaining lots):
  cabins   cabins and the shared subdivision work only, resort amenities deferred
  lean     Will's lean phase: 3,000 sf pavilion, 800 sf caterer's kitchen, plus the courts and trails
  full     Will's full program: 6,000 sf glass pavilion, 2,500 sf dining building, tennis + pickleball, trails

  python3 -I projects/montreve/scripts/model.py
"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
WP = "Will's business plan, Oct 22 2026"

# key, label, low, base, high, unit, source
INPUTS = [
    ('Program', None),
    ('existing', 'Financed homes (lots 1, 3, 6, 18, 21, 28, 30)', 7, 7, 7, 'homes', WP + ' p. 2'),
    ('new_cabins', 'New cabins on the remaining sale lots', 17, 17, 17, 'cabins', 'Site plan: 30 lots - 7 homes - 6 amenity lots'),
    ('cabin_sf', 'Cabin conditioned area (684 basement + 880 + 880)', 2444, 2444, 2444, 'sf', 'Owner cabin plans'),
    ('Spent to date (reported, unaudited)', None),
    ('sunk_land', 'Land', 1500000, 1500000, 1500000, '$', WP + ' p. 11 (Oct 20 cost workbook)'),
    ('sunk_gc', 'General contractor to date', 4512349.09, 4512349.09, 4512349.09, '$', WP + ' p. 11'),
    ('sunk_soft', 'Soft costs to date', 639581.60, 639581.60, 639581.60, '$', WP + ' p. 11'),
    ('sunk_int', 'Interest paid to date', 479850.11, 479850.11, 479850.11, '$', WP + ' p. 11'),
    ('Finish the 7 financed homes', None),
    ('existing_finish', 'Completion and turnover reserve (all 7)', 300000, 600000, 1000000, '$', WP + ' p. 13; GC shows $270,046 left on contract'),
    ('New cabins (per cabin)', None),
    ('cabin_hard', 'Vertical construction', 640000, 710000, 800000, '$/cabin', 'Project GC contract $4.78M / 7 = $683k, +4% escalation (base)'),
    ('cabin_site', 'Lot work: driveway, septic or sewer tap, power and water service', 40000, 65000, 120000, '$/cabin', 'Allowance; sloped lots, power is the known constraint (fractional plan)'),
    ('cabin_ffe', 'Furnishing to rental-ready (sold turnkey)', 40000, 55000, 75000, '$/cabin', 'Allowance'),
    ('cabin_soft', 'Soft costs (permits, engineering, builder\'s risk)', 0.05, 0.07, 0.09, '% of hard + lot', 'Allowance'),
    ('cabin_cont', 'Contingency (repeat plan)', 0.05, 0.08, 0.12, '% of hard + lot', 'Allowance'),
    ('Shared site work (every program)', None),
    ('predev', 'Survey, topo, geotech, civil, entitlements, HOA documents', 250000, 450000, 750000, '$', WP + ' p. 13'),
    ('infra', 'Roads to standard, utility mains, power extension, gate and entry, stormwater', 500000, 900000, 2000000, '$', 'Share of ' + WP + ' $0.8-3.0M site line'),
    ('Resort amenities: full program', None),
    ('pav_sf', 'Glass wedding pavilion', 6000, 6000, 6000, 'sf', WP + ' p. 6'),
    ('pav_psf', 'Pavilion cost', 400, 550, 700, '$/sf', WP + ' p. 13'),
    ('din_sf', 'Dining and catering building', 2500, 2500, 2500, 'sf', WP + ' p. 6'),
    ('din_psf', 'Dining cost', 300, 450, 600, '$/sf', WP + ' p. 13'),
    ('courts', 'One tennis and two pickleball courts', 150000, 220000, 350000, '$', WP + ' p. 13'),
    ('trails', 'Trails and landscape', 60000, 120000, 250000, '$', WP + ' p. 13'),
    ('amen_site', 'Event parking and amenity site work', 300000, 600000, 1000000, '$', 'Balance of ' + WP + ' site line'),
    ('amen_ffe', 'Amenity furniture and equipment', 200000, 300000, 500000, '$', WP + ' p. 13 (cabins excluded)'),
    ('preopen', 'Preopening and working capital', 250000, 400000, 600000, '$', WP + ' p. 13'),
    ('Resort amenities: lean program', None),
    ('lpav_sf', 'Smaller pavilion', 3000, 3000, 3000, 'sf', WP + ' p. 17'),
    ('lpav_psf', 'Smaller pavilion cost', 380, 500, 620, '$/sf', WP + ' p. 17'),
    ('lsvc_sf', 'Caterer service kitchen', 800, 800, 800, 'sf', WP + ' p. 17'),
    ('lsvc_psf', 'Service kitchen cost', 300, 400, 520, '$/sf', WP + ' p. 17'),
    ('lamen_site', 'Event parking and site work (lean)', 200000, 400000, 700000, '$', 'Allowance'),
    ('lamen_ffe', 'Amenity furniture (lean)', 100000, 150000, 250000, '$', 'Allowance'),
    ('lpreopen', 'Preopening (lean)', 150000, 250000, 400000, '$', WP + ' p. 17'),
    ('amen_cont', 'Contingency on amenity hard cost and furniture', 0.15, 0.20, 0.25, '%', WP + ' p. 13'),
    ('Financing and selling', None),
    ('rate', 'Construction loan interest', 0.08, 0.09, 0.105, '%/yr', WP + ' p. 16 uses 9%'),
    ('draw', 'Average share of new spend outstanding (sales repay as you go)', 0.35, 0.40, 0.50, '%', 'Phased build and sell'),
    ('months', 'Build-and-sell period', 18, 24, 30, 'months', WP + ' p. 19: 15-24+ months'),
    ('loan_fee', 'Loan fees', 0.0075, 0.01, 0.015, '% of new spend', 'Allowance'),
    ('debt', 'Existing loans (7 homes, Oct 14 2026)', 5303543.30, 5303543.30, 5303543.30, '$', WP + ' p. 12'),
    ('debt_months', 'Months the existing loans stay open', 6, 9, 14, 'months', 'Until the homes sell'),
    ('ext_fee', 'Extension fees per quarter (7 loans)', 59742.97, 59742.97, 59742.97, '$/qtr', WP + ' p. 12'),
    ('commission', 'Sales commission', 0.05, 0.06, 0.06, '% of price', 'Market'),
    ('closing', 'Seller closing, title, transfer tax', 0.01, 0.015, 0.02, '% of price', 'Market'),
    ('marketing', 'Sales marketing, model staging, website', 150000, 250000, 400000, '$', 'Allowance'),
    ('hold', 'Taxes, insurance and dues on unsold inventory', 100000, 150000, 250000, '$', 'Allowance'),
    ('Pricing tests', None),
    ('price', 'Cabin sale price tested', 1300000, 1300000, 1300000, '$/cabin', 'Matthew: break-even target'),
    ('margin', 'Developer margin target (on cost)', 0.15, 0.15, 0.15, '%', 'Typical for-sale target'),
    ('comp', 'Listed 3-bed indoor-pool cabins, Sevierville', 900000, 925000, 950000, '$', 'Active listings Oct 2026 (asking, not sold)'),
    ('Resort operations (stabilized year, full program)', None),
    ('weddings', 'Weddings per year', 20, 35, 50, 'events', WP + ' p. 14'),
    ('wed_fee', 'Venue fee per wedding', 4000, 6000, 8000, '$', WP + ' p. 14'),
    ('retreats', 'Retreats per year', 6, 10, 16, 'events', WP + ' p. 14'),
    ('ret_fee', 'Venue fee per retreat', 2500, 4000, 6000, '$', WP + ' p. 14'),
    ('ancillary', 'Kitchen use, bar and rentals per event', 500, 1000, 1500, '$', 'Allowance (catering itself is outsourced)'),
    ('other_rev', 'Other property revenue', 20000, 40000, 70000, '$', WP + ' p. 14'),
    ('pool_share', 'Cabins in the rental program', 0.60, 0.75, 0.85, '% of 24', 'Assumption'),
    ('gross_rent', 'Gross rent per cabin per year', 75000, 95000, 118625, '$', 'Comp $80k (2025); ' + WP + ' $650 x 50% = $118,625'),
    ('mgmt_fee', 'Rental management fee', 0.20, 0.25, 0.30, '% of rent', 'Local range 20-35%'),
    ('mgmt_cost', 'Direct cost of running the rental program', 0.55, 0.50, 0.45, '% of fee', 'Allowance'),
    ('dues', 'Owner club dues per cabin', 350, 450, 600, '$/month', 'Assumption; sets amenity funding'),
    ('venue_var', 'Event variable cost', 0.30, 0.30, 0.30, '% of event revenue', WP + ' p. 14'),
    ('other_var', 'Other revenue variable cost', 0.40, 0.40, 0.40, '%', WP + ' p. 14'),
    ('fixed', 'Fixed overhead (staff, insurance, utilities, grounds, marketing)', 480000, 480000, 480000, '$/yr', WP + ' p. 10'),
    ('reserve', 'Replacement reserve', 0.04, 0.04, 0.04, '% of revenue', WP + ' p. 14'),
    ('cap', 'Cap rate to value the resort business', 0.11, 0.10, 0.09, '%', 'Assumption'),
    ('Owner economics (per cabin)', None),
    ('own_util', 'Utilities, internet, pool and hot tub service', 9000, 10000, 12000, '$/yr', 'Allowance'),
    ('own_maint', 'Maintenance and supplies', 5000, 6000, 8000, '$/yr', 'Allowance'),
    ('own_ins', 'Insurance', 5000, 6000, 8000, '$/yr', 'Allowance'),
    ('own_tax', 'Property tax rate on sale price', 0.0040, 0.0045, 0.0050, '% of price/yr', 'Sevier Co. 25% assessment x ~$1.80/$100, unverified'),
]
ROW = {}   # key -> Inputs row
COLS = {'low': 'C', 'base': 'D', 'high': 'E'}
V = {s: {k: r[i] for r in INPUTS if r[1] for k in [r[0]]} for s, i in (('low', 2), ('base', 3), ('high', 4))}


def model(v):
    """The workbook's formulas, in Python (checked against the workbook by LibreOffice recalculation)."""
    units = v['existing'] + v['new_cabins']
    sunk = v['sunk_land'] + v['sunk_gc'] + v['sunk_soft'] + v['sunk_int']
    hl = v['cabin_hard'] + v['cabin_site']
    per_cabin = hl * (1 + v['cabin_soft'] + v['cabin_cont']) + v['cabin_ffe']
    cabins = per_cabin * v['new_cabins']
    shared = v['predev'] + v['infra']
    full_hard = v['pav_sf'] * v['pav_psf'] + v['din_sf'] * v['din_psf'] + v['courts'] + v['trails'] + v['amen_site'] + v['amen_ffe']
    lean_hard = v['lpav_sf'] * v['lpav_psf'] + v['lsvc_sf'] * v['lsvc_psf'] + v['courts'] + v['trails'] + v['lamen_site'] + v['lamen_ffe']
    amen = {'cabins': 0, 'lean': lean_hard * (1 + v['amen_cont']) + v['lpreopen'], 'full': full_hard * (1 + v['amen_cont']) + v['preopen']}
    carry_old = v['debt'] * v['rate'] * v['debt_months'] / 12 + v['ext_fee'] * v['debt_months'] / 3
    carry_k = v['rate'] * v['draw'] * v['months'] / 12 + v['loan_fee']
    sell_pct = v['commission'] + v['closing']
    out = {'units': units, 'sunk': sunk, 'per_cabin': per_cabin, 'per_cabin_sf': per_cabin / v['cabin_sf'], 'cabins': cabins, 'shared': shared,
           'existing_finish': v['existing_finish'], 'carry_old': carry_old, 'carry_k': carry_k, 'sell_pct': sell_pct, 'programs': {}}
    for p, a in amen.items():
        new_spend = v['existing_finish'] + cabins + shared + a
        carry = new_spend * carry_k
        fixed_sell = v['marketing'] + v['hold']
        cost = sunk + new_spend + carry + carry_old + fixed_sell
        be = cost / (units * (1 - sell_pct))
        need = cost * (1 + v['margin']) / (units * (1 - sell_pct))
        revenue = units * v['price']
        profit = revenue * (1 - sell_pct) - cost
        out['programs'][p] = {'amenities': a, 'new_spend': new_spend, 'carry_new': carry, 'cost': cost, 'break_even': be,
                              'price_for_margin': need, 'revenue_at_price': revenue, 'profit_at_price': profit,
                              'margin_at_price': profit / cost, 'premium_vs_comp': be / v['comp'] - 1}
        # the amenity budget a sale at the tested price can carry at break-even: solve cost(a) = net sales for a
        base_wo = sunk + v['existing_finish'] + cabins + shared
        out['programs'][p]['amenity_budget_at_price'] = (units * v['price'] * (1 - sell_pct) - carry_old - fixed_sell) / (1 + carry_k) - base_wo
    # resort operations, stabilized year (full program; the lean program runs on a $300k overhead, plan p. 17)
    events = v['weddings'] + v['retreats']
    venue = v['weddings'] * v['wed_fee'] + v['retreats'] * v['ret_fee']
    anc = events * v['ancillary']
    fee = units * v['pool_share'] * v['gross_rent'] * v['mgmt_fee']
    dues = units * v['dues'] * 12
    rev = venue + anc + v['other_rev'] + fee + dues
    costs = (venue + anc) * v['venue_var'] + v['other_rev'] * v['other_var'] + fee * v['mgmt_cost'] + v['fixed'] + rev * v['reserve']
    noi = rev - costs
    out['resort'] = {'events': events, 'venue': venue, 'ancillary': anc, 'other': v['other_rev'], 'mgmt_fees': fee, 'dues': dues, 'revenue': rev,
                     'event_var': (venue + anc) * v['venue_var'], 'other_var': v['other_rev'] * v['other_var'], 'mgmt_cost': fee * v['mgmt_cost'],
                     'fixed': v['fixed'], 'reserve': rev * v['reserve'], 'costs': costs, 'noi': noi, 'value': max(0, noi) / v['cap'],
                     'dues_breakeven': None}
    # monthly dues that would cover the resort's shortfall (or 0)
    out['resort']['dues_breakeven'] = max(0, v['dues'] - noi / (units * 12 * (1 - v['reserve'])))
    # owner: only the rent changes across the columns; fee and costs stay at base so the columns compare like for like
    b = V['base']; g = v['gross_rent']; mg = g * b['mgmt_fee']; tax = b['price'] * b['own_tax']; d = b['dues'] * 12
    exp = mg + b['own_util'] + b['own_maint'] + b['own_ins'] + tax + d
    v = {**v, **{k: b[k] for k in ('own_util', 'own_maint', 'own_ins')}}
    out['owner'] = {'gross': g, 'mgmt': mg, 'util': v['own_util'], 'maint': v['own_maint'], 'ins': v['own_ins'], 'tax': tax, 'dues': d,
                    'expenses': exp, 'net': g - exp, 'yield': (g - exp) / v['price']}
    return out


# ---------------------------------------------------------------- workbook
NAVY, PALE, LINE = '0A3047', 'EAF1F6', 'C9D6E0'
B = Font(bold=True); H = Font(bold=True, color='FFFFFF'); T = Font(bold=True, size=14, color=NAVY)
FILL_H, FILL_S, FILL_IN = PatternFill('solid', fgColor=NAVY), PatternFill('solid', fgColor=PALE), PatternFill('solid', fgColor='FFF7DD')
thin = Border(bottom=Side(style='thin', color=LINE))
MONEY, PCT, NUM = '$#,##0;[Red]-$#,##0', '0.0%', '#,##0'


def fmt_for(unit):
    return PCT if unit.startswith('%') else (NUM if unit in ('homes', 'cabins', 'sf', 'events', 'months') else MONEY)


def head(ws, r, labels, widths=None):
    for i, t in enumerate(labels):
        c = ws.cell(r, i + 1, t); c.font = H; c.fill = FILL_H; c.alignment = Alignment(horizontal='center' if i else 'left', wrap_text=True)
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w


def inputs_sheet(wb):
    ws = wb.active; ws.title = 'Inputs'
    ws['A1'] = 'Montreve: model inputs'; ws['A1'].font = T
    ws['A2'] = 'Yellow cells are inputs. Every other sheet is formulas on these. Low / base / high are cost-side ranges; for revenue lines they are downside / base / upside.'
    head(ws, 4, ['Key', 'Input', 'Low', 'Base', 'High', 'Unit', 'Source'], [12, 58, 14, 14, 14, 16, 60])
    r = 5
    for row in INPUTS:
        if row[1] is None:
            c = ws.cell(r, 1, row[0]); c.font = B; ws.cell(r, 1).fill = FILL_S
            for k in range(2, 8): ws.cell(r, k).fill = FILL_S
            r += 1; continue
        k, label, lo, ba, hi, unit, src = row
        ROW[k] = r
        ws.cell(r, 1, k).font = Font(color='888888', size=9)
        ws.cell(r, 2, label)
        for col, val in zip('CDE', (lo, ba, hi)):
            c = ws[f'{col}{r}']; c.value = val; c.number_format = fmt_for(unit); c.fill = FILL_IN
        ws.cell(r, 6, unit); ws.cell(r, 7, src).font = Font(size=9, color='555555')
        r += 1
    ws.freeze_panes = 'C5'


def I(k, col):  # an input cell
    return f"Inputs!${col}${ROW[k]}"


def budget_sheet(wb):
    ws = wb.create_sheet('Budget')
    ws['A1'] = 'Development budget and break-even'; ws['A1'].font = T
    ws['A2'] = 'All 24 cabins for sale; the amenity line changes with the program. Columns: low / base / high cost case.'
    head(ws, 4, ['Line', 'Low', 'Base', 'High', 'Note'], [62, 16, 16, 16, 60])
    lines, r = {}, 5

    def put(key, label, f, note='', fmt=MONEY, bold=False, shade=False):
        nonlocal r
        ws.cell(r, 1, label).font = B if bold else Font()
        for j, col in enumerate('CDE'):
            c = ws.cell(r, 2 + j, '=' + f(col)); c.number_format = fmt
            if bold: c.font = B
        ws.cell(r, 5, note).font = Font(size=9, color='555555')
        if shade:
            for k in range(1, 6): ws.cell(r, k).fill = FILL_S
        lines[key] = r; r += 1

    L = lambda key, col: f"{'BCD'['CDE'.index(col)]}{lines[key]}"   # a budget cell in the same case column
    put('units', 'Cabins for sale', lambda c: f"{I('existing', c)}+{I('new_cabins', c)}", fmt=NUM)
    put('sunk', 'Spent to date (land, GC, soft, interest; reported)', lambda c: '+'.join(I(k, c) for k in ('sunk_land', 'sunk_gc', 'sunk_soft', 'sunk_int')), 'Oct 20 workbook; not yet reconciled to the ledger')
    put('fin', 'Finish the 7 financed homes', lambda c: I('existing_finish', c))
    put('per', 'One new cabin, all in', lambda c: f"({I('cabin_hard', c)}+{I('cabin_site', c)})*(1+{I('cabin_soft', c)}+{I('cabin_cont', c)})+{I('cabin_ffe', c)}", 'Hard + lot work, soft costs, contingency, furniture')
    put('persf', '   per sf', lambda c: f"{L('per', c)}/{I('cabin_sf', c)}")
    put('cab', 'New cabins (17)', lambda c: f"{L('per', c)}*{I('new_cabins', c)}")
    put('shared', 'Shared site work and predevelopment', lambda c: f"{I('predev', c)}+{I('infra', c)}")
    put('lean', 'Resort amenities: lean program', lambda c: f"({I('lpav_sf', c)}*{I('lpav_psf', c)}+{I('lsvc_sf', c)}*{I('lsvc_psf', c)}+{I('courts', c)}+{I('trails', c)}+{I('lamen_site', c)}+{I('lamen_ffe', c)})*(1+{I('amen_cont', c)})+{I('lpreopen', c)}", '3,000 sf pavilion, 800 sf service kitchen, courts, trails')
    put('full', 'Resort amenities: full program', lambda c: f"({I('pav_sf', c)}*{I('pav_psf', c)}+{I('din_sf', c)}*{I('din_psf', c)}+{I('courts', c)}+{I('trails', c)}+{I('amen_site', c)}+{I('amen_ffe', c)})*(1+{I('amen_cont', c)})+{I('preopen', c)}", '6,000 sf pavilion, 2,500 sf dining, courts, trails')
    put('ck', 'Carry on new spend (share of it)', lambda c: f"{I('rate', c)}*{I('draw', c)}*{I('months', c)}/12+{I('loan_fee', c)}", 'Interest on the average balance + loan fees', fmt=PCT)
    put('old', 'Carry on the existing loans until the homes sell', lambda c: f"{I('debt', c)}*{I('rate', c)}*{I('debt_months', c)}/12+{I('ext_fee', c)}*{I('debt_months', c)}/3")
    put('fsell', 'Marketing and unsold-inventory holding', lambda c: f"{I('marketing', c)}+{I('hold', c)}")
    put('sp', 'Selling cost (commission + closing)', lambda c: f"{I('commission', c)}+{I('closing', c)}", fmt=PCT)
    r += 1
    for p, label, amen in (('cabins', 'Program A: cabins only (amenities deferred)', None), ('lean', 'Program B: cabins + lean resort', 'lean'), ('full', 'Program C: cabins + full resort', 'full')):
        ws.cell(r, 1, label).font = Font(bold=True, color=NAVY); r += 1
        a = (lambda c: '0') if not amen else (lambda c, amen=amen: L(amen, c))
        put(p + '_new', '   New spend', lambda c, a=a: f"{L('fin', c)}+{L('cab', c)}+{L('shared', c)}+{a(c)}")
        put(p + '_cost', '   Total cost (to date + new + financing + selling fixed)', lambda c, p=p: f"{L('sunk', c)}+{L(p + '_new', c)}*(1+{L('ck', c)})+{L('old', c)}+{L('fsell', c)}", bold=True)
        put(p + '_be', '   Break-even price per cabin', lambda c, p=p: f"{L(p + '_cost', c)}/({L('units', c)}*(1-{L('sp', c)}))", bold=True, shade=True)
        put(p + '_m', '   Price for the margin target', lambda c, p=p: f"{L(p + '_cost', c)}*(1+{I('margin', c)})/({L('units', c)}*(1-{L('sp', c)}))")
        put(p + '_pr', '   Profit if every cabin sells at the tested price', lambda c, p=p: f"{L('units', c)}*{I('price', c)}*(1-{L('sp', c)})-{L(p + '_cost', c)}")
        put(p + '_prem', '   Break-even vs listed comparable cabins', lambda c, p=p: f"{L(p + '_be', c)}/{I('comp', c)}-1", fmt=PCT)
        r += 1
    ws.cell(r, 1, 'Amenity budget a sale at the tested price can carry (break-even)').font = B
    for j, c in enumerate('CDE'):
        f = f"=({L('units', c)}*{I('price', c)}*(1-{L('sp', c)})-{L('old', c)}-{L('fsell', c)})/(1+{L('ck', c)})-{L('sunk', c)}-{L('fin', c)}-{L('cab', c)}-{L('shared', c)}"
        x = ws.cell(r, 2 + j, f); x.number_format = MONEY; x.font = B
    ws.cell(r, 5, 'Negative: the tested price does not cover the cabins and site work before any amenity').font = Font(size=9, color='555555')
    lines['amen_cap'] = r
    ws.freeze_panes = 'B5'
    return lines


def resort_sheet(wb):
    ws = wb.create_sheet('Resort operations')
    ws['A1'] = 'Resort operations: stabilized year (full program)'; ws['A1'].font = T
    ws['A2'] = 'Venue, rental management and club dues. Catering is outsourced: food spend is not resort revenue. Downside / base / upside.'
    head(ws, 4, ['Line', 'Downside', 'Base', 'Upside', 'How'], [52, 16, 16, 16, 60])
    r, L = 5, {}

    def put(key, label, f, note='', fmt=MONEY, bold=False):
        nonlocal r
        ws.cell(r, 1, label).font = B if bold else Font()
        for j, col in enumerate('CDE'):
            c = ws.cell(r, 2 + j, '=' + f(col)); c.number_format = fmt
            if bold: c.font = B
        ws.cell(r, 5, note).font = Font(size=9, color='555555'); L[key] = r; r += 1
    R = lambda key, col: f"{'BCD'['CDE'.index(col)]}{L[key]}"
    units = lambda c: f"({I('existing', c)}+{I('new_cabins', c)})"
    put('venue', 'Venue fees (weddings + retreats)', lambda c: f"{I('weddings', c)}*{I('wed_fee', c)}+{I('retreats', c)}*{I('ret_fee', c)}")
    put('anc', 'Kitchen use, bar and rentals', lambda c: f"({I('weddings', c)}+{I('retreats', c)})*{I('ancillary', c)}")
    put('oth', 'Other property revenue', lambda c: I('other_rev', c))
    put('fee', 'Rental management fees', lambda c: f"{units(c)}*{I('pool_share', c)}*{I('gross_rent', c)}*{I('mgmt_fee', c)}", 'Cabins in program x gross rent x fee')
    put('dues', 'Owner club dues', lambda c: f"{units(c)}*{I('dues', c)}*12")
    put('rev', 'Revenue', lambda c: f"SUM({R('venue', c)}:{R('dues', c)})", bold=True)
    put('ev', 'Event variable cost', lambda c: f"({R('venue', c)}+{R('anc', c)})*{I('venue_var', c)}")
    put('ov', 'Other variable cost', lambda c: f"{R('oth', c)}*{I('other_var', c)}")
    put('mc', 'Rental program direct cost', lambda c: f"{R('fee', c)}*{I('mgmt_cost', c)}")
    put('fx', 'Fixed overhead', lambda c: I('fixed', c))
    put('rs', 'Replacement reserve', lambda c: f"{R('rev', c)}*{I('reserve', c)}")
    put('cost', 'Operating cost', lambda c: f"SUM({R('ev', c)}:{R('rs', c)})", bold=True)
    put('noi', 'Net operating income', lambda c: f"{R('rev', c)}-{R('cost', c)}", bold=True)
    put('val', 'Value of the resort business at the cap rate', lambda c: f"MAX(0,{R('noi', c)})/{I('cap', c)}")
    return L


def owner_sheet(wb):
    ws = wb.create_sheet('Owner economics')
    ws['A1'] = 'What a cabin owner can expect: one cabin, a year in the rental program'; ws['A1'].font = T
    ws['A2'] = 'Before any mortgage and income tax. Illustration for discussion, not a promise of income.'
    head(ws, 4, ['Line', 'Downside', 'Base', 'Upside', ''], [52, 16, 16, 16, 30])
    r, L = 5, {}
    rows = [('g', 'Gross rent', lambda c: I('gross_rent', c)),
            ('m', 'Management fee', lambda c: f"-{I('gross_rent', c)}*{I('mgmt_fee', 'D')}"),
            ('u', 'Utilities, internet, pool service', lambda c: f"-{I('own_util', 'D')}"),
            ('mt', 'Maintenance and supplies', lambda c: f"-{I('own_maint', 'D')}"),
            ('in', 'Insurance', lambda c: f"-{I('own_ins', 'D')}"),
            ('t', 'Property tax', lambda c: f"-{I('price', 'D')}*{I('own_tax', 'D')}"),
            ('d', 'Club dues', lambda c: f"-{I('dues', 'D')}*12")]
    for k, label, f in rows:
        ws.cell(r, 1, label)
        for j, c in enumerate('CDE'): x = ws.cell(r, 2 + j, '=' + f(c)); x.number_format = MONEY
        L[k] = r; r += 1
    ws.cell(r, 1, 'Net to owner').font = B
    for j, c in enumerate('BCD'): x = ws.cell(r, 2 + j, f'=SUM({c}{L["g"]}:{c}{L["d"]})'); x.number_format = MONEY; x.font = B
    L['net'] = r; r += 1
    ws.cell(r, 1, 'Cash yield on the purchase price').font = B
    for j, (c, ic) in enumerate(zip('BCD', 'CDE')): x = ws.cell(r, 2 + j, f'={c}{L["net"]}/{I("price", ic)}'); x.number_format = PCT; x.font = B
    L['yield'] = r
    ws.cell(r + 2, 1, 'Only the gross rent changes across the columns; the fee, costs and dues are the base inputs.').font = Font(size=9, color='555555')
    return L


def sensitivity_sheet(wb, bl):
    ws = wb.create_sheet('Price sensitivity')
    ws['A1'] = 'Profit at each sale price (base costs)'; ws['A1'].font = T
    head(ws, 3, ['Price per cabin', 'A: cabins only', 'B: lean resort', 'C: full resort'], [20, 20, 20, 20])
    for i, price in enumerate(range(1000000, 1700001, 50000)):
        r = 4 + i
        ws.cell(r, 1, price).number_format = MONEY
        for j, p in enumerate(('cabins', 'lean', 'full')):
            x = ws.cell(r, 2 + j, f"=Budget!$C${bl['units']}*$A{r}*(1-Budget!$C${bl['sp']})-Budget!$C${bl[p + '_cost']}"); x.number_format = MONEY


def cashflow_sheet(wb, bl):
    """Quarterly base-case cash flow for the program chosen in B3 (A, B or C). Spend and sales timing are inputs."""
    ws = wb.create_sheet('Cash flow')
    ws['A1'] = 'Quarterly cash flow, base costs, at the tested price'; ws['A1'].font = T
    ws['A3'] = 'Program (A, B or C):'; ws['B3'] = 'C'; ws['B3'].fill = FILL_IN
    ws['A4'] = 'Timing rows (yellow) are the share of each cost spent, or cabins sold, in each quarter. Existing loans are repaid from the first home sales.'
    qs = [f'{y} Q{q}' for y in range(2027, 2031) for q in range(1, 5)]
    head(ws, 6, ['Line'] + qs + ['Total'], [44] + [11] * len(qs) + [13])
    timing = {
        'fin': [.5, .5] + [0] * 14,
        'shared': [.2, .3, .3, .2] + [0] * 12,
        'amen': [0, .1, .2, .25, .25, .2] + [0] * 10,
        'cab': [0, .06, .12, .12, .12, .12, .12, .12, .12, .1] + [0] * 6,
        'sell_existing': [2, 3, 2] + [0] * 13,
        'sell_new': [0, 0, 0, 1, 2, 2, 2, 2, 2, 2, 2, 2] + [0] * 4,
    }
    names = {'fin': 'Finish 7 homes (share)', 'shared': 'Shared site work (share)', 'amen': 'Amenities (share)', 'cab': 'New cabins (share)',
             'sell_existing': 'Financed homes sold (count)', 'sell_new': 'New cabins sold (count)'}
    r, rows = 7, {}
    for k, vals in timing.items():
        ws.cell(r, 1, names[k])
        for j, v in enumerate(vals): c = ws.cell(r, 2 + j, v); c.fill = FILL_IN; c.number_format = '0%' if 'share' in names[k] else '0'
        ws.cell(r, 18, f'=SUM(B{r}:Q{r})').number_format = '0%' if 'share' in names[k] else '0'
        rows[k] = r; r += 1
    r += 1
    amen = f'IF($B$3="A",0,IF($B$3="B",Budget!$C${bl["lean"]},Budget!$C${bl["full"]}))'
    lines = [
        ('spend', 'Development spend', lambda c: f"-({c}{rows['fin']}*Budget!$C${bl['fin']}+{c}{rows['shared']}*Budget!$C${bl['shared']}+{c}{rows['amen']}*{amen}+{c}{rows['cab']}*Budget!$C${bl['cab']})"),
        ('fin_cost', 'Financing (spread with spend)', lambda c: f"{c}{{spend}}*Budget!$C${bl['ck']}"),
        ('old', 'Carry on existing loans (while open)', lambda c: f"-IF(SUM($B{rows['sell_existing']}:{c}{rows['sell_existing']})-{c}{rows['sell_existing']}<{I('existing', 'D')},Budget!$C${bl['old']}/ROUNDUP({I('debt_months', 'D')}/3,0),0)"),
        ('fixed', 'Marketing and holding', lambda c: f"-Budget!$C${bl['fsell']}/16"),
        ('repay', 'Repay existing loans (as each financed home sells)', lambda c: f"-{c}{rows['sell_existing']}/{I('existing', 'D')}*{I('debt', 'D')}"),
        ('sales', 'Net sales proceeds', lambda c: f"({c}{rows['sell_existing']}+{c}{rows['sell_new']})*{I('price', 'D')}*(1-Budget!$C${bl['sp']})"),
        ('net', 'Net cash flow', lambda c: f"SUM({c}{{spend}}:{c}{{sales}})"),
        ('cum', 'Cumulative', lambda c: f"{c}{{net}}" if c == 'B' else f"{chr(ord(c) - 1)}{{cum}}+{c}{{net}}"),
    ]
    at = {}
    for k, label, _ in lines: at[k] = r; r += 1
    for k, label, f in lines:
        rr = at[k]; ws.cell(rr, 1, label).font = B if k in ('net', 'cum') else Font()
        for j in range(16):
            c = get_column_letter(2 + j)
            x = ws.cell(rr, 2 + j, '=' + f(c).format(**at)); x.number_format = MONEY
        if k != 'cum':
            ws.cell(rr, 18, f'=SUM(B{rr}:Q{rr})').number_format = MONEY
    r += 1
    ws.cell(r, 1, 'Peak cash need (most negative cumulative)').font = B
    x = ws.cell(r, 2, f"=MIN(B{at['cum']}:Q{at['cum']})"); x.number_format = MONEY; x.font = B
    ws.cell(r + 1, 1, 'Spent to date is not in this sheet (already spent). Peak cash need is what equity and construction loans must cover together.').font = Font(size=9, color='555555')
    ws.freeze_panes = 'B7'
    return timing


def main():
    res = {s: model(V[s]) for s in V}
    wb = Workbook(); inputs_sheet(wb)
    bl = budget_sheet(wb); rl = resort_sheet(wb); ol = owner_sheet(wb); sensitivity_sheet(wb, bl); timing = cashflow_sheet(wb, bl)
    out = ROOT / 'out' / 'Montreve_Budget_ProForma.xlsx'; out.parent.mkdir(exist_ok=True); wb.save(out)
    (ROOT / 'data' / 'model.json').write_text(json.dumps({'inputs': V, 'results': res, 'cells': {'budget': bl, 'resort': rl, 'owner': ol}, 'timing': timing}, indent=1, default=float))
    b = res['base']
    print('per new cabin $%.0f (%.0f/sf)' % (b['per_cabin'], b['per_cabin_sf']))
    for p, x in b['programs'].items():
        print(f"{p:7s} cost ${x['cost'] / 1e6:6.2f}M  break-even ${x['break_even'] / 1e3:,.0f}k  (low ${res['low']['programs'][p]['break_even'] / 1e3:,.0f}k high ${res['high']['programs'][p]['break_even'] / 1e3:,.0f}k)"
              f"  profit@1.3M ${x['profit_at_price'] / 1e6:5.2f}M  amenity budget@1.3M ${x['amenity_budget_at_price'] / 1e6:5.2f}M")
    print('resort NOI', {s: round(res[s]['resort']['noi']) for s in res}, 'owner net', {s: round(res[s]['owner']['net']) for s in res})
    print('->', out)


if __name__ == '__main__':
    main()
